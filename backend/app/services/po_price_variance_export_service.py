"""
Excel export for PR vs PO price variance report.
"""
from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from typing import Optional
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.schemas.procurement import PrPoPriceVarianceRow
from app.services.excel_export_utils import (
    apply_merged_row_style,
    configure_print_layout,
    date_range_filename_suffix,
    finalize_column_widths,
    format_date_range_label,
    row_height_for_text,
    set_row_heights_for_wrapped_rows,
)

WIB = ZoneInfo("Asia/Jakarta")

QTY_FORMAT_INT = "#,##0"
QTY_FORMAT_DEC = "#,##0.####"
IDR_FORMAT = '"Rp" #,##0;"Rp" -#,##0'

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
TITLE_FONT = Font(name="Calibri", size=18, bold=True, color="1F4E79")
SUBTITLE_FONT = Font(name="Calibri", size=11, color="6B7280")
META_LABEL_FONT = Font(name="Calibri", size=10, bold=True, color="374151")
META_VALUE_FONT = Font(name="Calibri", size=10, color="111827")
DATA_FONT = Font(name="Calibri", size=10, color="111827")
TOTAL_FONT = Font(name="Calibri", size=10, bold=True, color="111827")
TOTAL_FILL = PatternFill("solid", fgColor="E5E7EB")
META_FILL = PatternFill("solid", fgColor="F8FAFC")
ALT_ROW_FILL = PatternFill("solid", fgColor="F9FAFB")

THIN = Side(style="thin", color="D1D5DB")
MEDIUM = Side(style="medium", color="9CA3AF")

THIN_BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
TOTAL_TOP_BORDER = Border(left=THIN, right=THIN, top=MEDIUM, bottom=THIN)

# Header multi-baris agar muat tanpa terpotong di layar maupun cetak.
COLUMNS = [
    ("No.", 6),
    ("No. PO", 18),
    ("No. PR", 18),
    ("Nama Produk", 34),
    ("Unit\nPemohon", 18),
    ("Vendor", 26),
    ("Qty", 10),
    ("Harga Satuan\nPR (Rp)", 18),
    ("Harga Satuan\nPO Awal (Rp)", 18),
    ("Harga Satuan\nPO Saat Ini (Rp)", 20),
    ("Selisih Satuan\nvs PR (Rp)", 18),
    ("Total PO\nSaat Ini (Rp)", 18),
    ("Selisih Total\nvs PR (Rp)", 18),
    ("Status\nPembayaran", 16),
    ("Keterangan", 22),
]

TOTAL_COLS = (12, 13)
QTY_COL = 7
IDR_MIN_WIDTH = 20
WRAP_TEXT_COLS = (4, 5, 6, 15)


def _decimal(value) -> Decimal:
    if value is None:
        return Decimal("0")
    return Decimal(str(value))


def _is_whole_qty(value: Decimal) -> bool:
    normalized = value.quantize(Decimal("0.0001"))
    return normalized == normalized.to_integral_value()


def _qty_number_format(value: Decimal) -> str:
    return QTY_FORMAT_INT if _is_whole_qty(value) else QTY_FORMAT_DEC


def _qty_value(value: Decimal) -> float | int:
    normalized = value.quantize(Decimal("0.0001"))
    if normalized == normalized.to_integral_value():
        return int(normalized)
    return float(normalized)


def _idr_value(value: Decimal) -> int:
    return int(value.quantize(Decimal("1")))


def _idr_display_len(value: Decimal) -> int:
    amount = _idr_value(value)
    return len(f"Rp {amount:,}".replace(",", "."))


def _notes_for_row(row: PrPoPriceVarianceRow) -> str:
    notes: list[str] = []
    if row.has_price_change:
        notes.append("Revisi PO")
    if row.differs_from_pr:
        notes.append("Beda Harga PR")
    return ", ".join(notes) if notes else "-"


def _filter_label(
    *,
    search: Optional[str],
    unit_name: Optional[str],
    only_variance: bool,
    unpaid_only: bool,
    date_from: date | None = None,
    date_to: date | None = None,
) -> list[tuple[str, str]]:
    return [
        ("Periode Laporan", format_date_range_label(date_from, date_to)),
        ("Pencarian", search or "—"),
        ("Unit Pemohon", unit_name or "Semua Unit"),
        ("Kriteria", _criteria_label(only_variance, unpaid_only)),
        ("Waktu Cetak", datetime.now(WIB).strftime("%d/%m/%Y %H:%M WIB")),
    ]


def _criteria_label(only_variance: bool, unpaid_only: bool) -> str:
    parts: list[str] = []
    if only_variance:
        parts.append("Hanya Selisih Harga")
    if unpaid_only:
        parts.append("Hanya Belum Lunas")
    return ", ".join(parts) if parts else "Semua Data"


def price_variance_filename(
    date_from: date | None = None,
    date_to: date | None = None,
) -> str:
    stamp = datetime.now(WIB).strftime("%Y%m%d")
    return f"laporan-selisih-harga-pr-po{date_range_filename_suffix(date_from, date_to)}_{stamp}.xlsx"


def generate_price_variance_excel(
    rows: list[PrPoPriceVarianceRow],
    *,
    generated_by: Optional[str] = None,
    search: Optional[str] = None,
    unit_name: Optional[str] = None,
    only_variance: bool = False,
    unpaid_only: bool = False,
    date_from: date | None = None,
    date_to: date | None = None,
) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Selisih Harga PR vs PO"

    last_col = len(COLUMNS)
    last_col_letter = get_column_letter(last_col)

    ws.merge_cells(f"A1:{last_col_letter}1")
    title = ws["A1"]
    title.value = "LAPORAN SELISIH HARGA PR vs PO"
    title.font = TITLE_FONT
    title.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 34

    ws.merge_cells(f"A2:{last_col_letter}2")
    subtitle = ws["A2"]
    subtitle.value = "Purchasing Go — Audit Perbandingan Harga Pengajuan dan Realisasi Pembelian"
    subtitle.font = SUBTITLE_FONT
    subtitle.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[2].height = row_height_for_text(subtitle.value, chars_per_line=90, min_height=24)

    meta_rows = _filter_label(
        search=search,
        unit_name=unit_name,
        only_variance=only_variance,
        unpaid_only=unpaid_only,
        date_from=date_from,
        date_to=date_to,
    )
    if generated_by:
        meta_rows.append(("Petugas / User", generated_by))
    meta_rows.append(("Jumlah Baris", str(len(rows))))

    meta_start = 4
    meta_end = meta_start + len(meta_rows) - 1

    for offset, (label, value) in enumerate(meta_rows):
        row_num = meta_start + offset
        
        # Merge A-C for the label
        ws.merge_cells(start_row=row_num, start_column=1, end_row=row_num, end_column=3)
        label_text = f"{label}:" if not label.endswith(":") else label
        label_cell = ws.cell(row=row_num, column=1, value=label_text)
        label_cell.font = META_LABEL_FONT
        label_cell.alignment = Alignment(horizontal="left", vertical="center")

        # Merge D-last_col for the value
        ws.merge_cells(start_row=row_num, start_column=4, end_row=row_num, end_column=last_col)
        value_cell = ws.cell(row=row_num, column=4, value=value)
        value_cell.font = META_VALUE_FONT
        value_cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        
        apply_merged_row_style(
            ws,
            row_num,
            1,
            last_col,
            fill=None,
            border=None,
        )
        ws.row_dimensions[row_num].height = row_height_for_text(
            value,
            chars_per_line=80,
            min_height=20,
        )

    header_row = meta_end + 2
    for col_idx, (header, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=header_row, column=col_idx, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.border = THIN_BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    ws.row_dimensions[header_row].height = 46

    data_start = header_row + 1

    if not rows:
        ws.merge_cells(start_row=data_start, start_column=1, end_row=data_start, end_column=last_col)
        note = ws.cell(row=data_start, column=1, value="Tidak ada data sesuai filter yang dipilih.")
        note.font = Font(name="Calibri", size=10, italic=True, color="6B7280")
        note.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        apply_merged_row_style(ws, data_start, 1, last_col, border=THIN_BORDER)
        ws.row_dimensions[data_start].height = 28
        total_row = data_start + 1
        last_data_row = data_start
    else:
        for index, item in enumerate(rows, start=1):
            row_num = data_start + index - 1
            row_fill = ALT_ROW_FILL if index % 2 == 0 else None
            qty = _decimal(item.real_qty)

            text_values = [
                index,
                item.po_number,
                item.pr_number or "-",
                item.product_name,
                item.unit_name or "-",
                item.vendor_name or "-",
            ]
            alignments = ["center", "left", "left", "left", "left", "left"]
            for col_idx, (text, align) in enumerate(zip(text_values, alignments), start=1):
                cell = ws.cell(row=row_num, column=col_idx, value=text)
                cell.font = DATA_FONT
                cell.alignment = Alignment(
                    horizontal=align,
                    vertical="center",
                    wrap_text=col_idx in WRAP_TEXT_COLS or col_idx in (2, 3),
                )
                cell.border = THIN_BORDER
                if row_fill:
                    cell.fill = row_fill

            qty_cell = ws.cell(row=row_num, column=QTY_COL, value=_qty_value(qty))
            qty_cell.font = DATA_FONT
            qty_cell.alignment = Alignment(horizontal="right", vertical="center")
            qty_cell.number_format = _qty_number_format(qty)
            qty_cell.border = THIN_BORDER
            if row_fill:
                qty_cell.fill = row_fill

            idr_fields = [
                item.pr_request_price,
                item.original_po_price,
                item.current_po_price,
                item.price_variance,
                item.current_po_total,
                item.total_variance,
            ]
            for offset, amount in enumerate(idr_fields, start=8):
                cell = ws.cell(row=row_num, column=offset, value=_idr_value(_decimal(amount)))
                cell.font = DATA_FONT
                cell.alignment = Alignment(horizontal="right", vertical="center")
                cell.number_format = IDR_FORMAT
                cell.border = THIN_BORDER
                if row_fill:
                    cell.fill = row_fill

            status_cell = ws.cell(row=row_num, column=14, value=item.payment_status or "-")
            status_cell.font = DATA_FONT
            status_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            status_cell.border = THIN_BORDER
            if row_fill:
                status_cell.fill = row_fill

            notes_cell = ws.cell(row=row_num, column=15, value=_notes_for_row(item))
            notes_cell.font = DATA_FONT
            notes_cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            notes_cell.border = THIN_BORDER
            if row_fill:
                notes_cell.fill = row_fill

            wrapped = [(col, COLUMNS[col - 1][1]) for col in WRAP_TEXT_COLS]
            wrapped.extend([(2, COLUMNS[1][1]), (3, COLUMNS[2][1]), (14, COLUMNS[13][1])])
            set_row_heights_for_wrapped_rows(ws, row_num, wrapped, min_height=22)

        total_row = data_start + len(rows)
        last_data_row = total_row - 1

    ws.merge_cells(start_row=total_row, start_column=1, end_row=total_row, end_column=11)
    total_label = ws.cell(row=total_row, column=1, value="JUMLAH")
    total_label.font = TOTAL_FONT
    total_label.fill = TOTAL_FILL
    total_label.alignment = Alignment(horizontal="right", vertical="center", indent=1)
    total_label.border = TOTAL_TOP_BORDER
    apply_merged_row_style(ws, total_row, 1, 11, fill=TOTAL_FILL, border=TOTAL_TOP_BORDER)

    if rows:
        first_data_row = data_start
        for col in TOTAL_COLS:
            col_letter = get_column_letter(col)
            cell = ws.cell(
                row=total_row,
                column=col,
                value=f"=SUM({col_letter}{first_data_row}:{col_letter}{last_data_row})",
            )
            cell.font = TOTAL_FONT
            cell.fill = TOTAL_FILL
            cell.border = TOTAL_TOP_BORDER
            cell.alignment = Alignment(horizontal="right", vertical="center")
            cell.number_format = IDR_FORMAT
        for col in range(14, last_col + 1):
            cell = ws.cell(row=total_row, column=col, value="")
            cell.fill = TOTAL_FILL
            cell.border = TOTAL_TOP_BORDER
    else:
        for col in range(12, last_col + 1):
            cell = ws.cell(row=total_row, column=col, value="" if col > 13 else 0)
            cell.font = TOTAL_FONT
            cell.fill = TOTAL_FILL
            cell.border = TOTAL_TOP_BORDER
            if col in TOTAL_COLS:
                cell.number_format = IDR_FORMAT
                cell.alignment = Alignment(horizontal="right", vertical="center")

    ws.row_dimensions[total_row].height = 24

    if rows:
        idr_fields_map = {
            8: "pr_request_price",
            9: "original_po_price",
            10: "current_po_price",
            11: "price_variance",
            12: "current_po_total",
            13: "total_variance",
        }
        idr_overrides = {}
        for col_idx, field in idr_fields_map.items():
            max_len = max(
                (_idr_display_len(_decimal(getattr(item, field))) for item in rows),
                default=IDR_MIN_WIDTH,
            )
            idr_overrides[col_idx] = max(IDR_MIN_WIDTH, max_len + 2)

        finalize_column_widths(
            ws,
            COLUMNS,
            header_row=header_row,
            data_start=data_start,
            data_end=last_data_row,
            overrides=idr_overrides,
        )

    footer_row = total_row + 2
    ws.merge_cells(start_row=footer_row, start_column=1, end_row=footer_row + 1, end_column=last_col)
    footer = ws.cell(
        row=footer_row,
        column=1,
        value=(
            "Catatan: Selisih dihitung dari Harga Satuan/Total PO Saat Ini dikurangi nilai PR. "
            "Dokumen ini dihasilkan otomatis oleh sistem Purchasing Go dan sah tanpa tanda tangan basah."
        ),
    )
    footer.font = Font(name="Calibri", size=9, italic=True, color="9CA3AF")
    footer.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
    ws.row_dimensions[footer_row].height = row_height_for_text(
        footer.value,
        chars_per_line=110,
        min_height=40,
        max_height=80,
    )
    ws.row_dimensions[footer_row + 1].height = 8

    ws.freeze_panes = ws.cell(row=data_start, column=1).coordinate
    ws.sheet_view.showGridLines = True
    configure_print_layout(ws, header_row=header_row, landscape=True, paper_size=8)

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
