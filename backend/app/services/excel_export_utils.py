"""
Shared helpers for professional Excel report exports.
"""
from __future__ import annotations

from datetime import date

from openpyxl.styles import Alignment, Border, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

_ID_MONTHS = (
    "Januari", "Februari", "Maret", "April", "Mei", "Juni",
    "Juli", "Agustus", "September", "Oktober", "November", "Desember",
)


def format_report_date(value: date) -> str:
    return f"{value.day:02d} {_ID_MONTHS[value.month - 1]} {value.year}"


def format_date_range_label(
    date_from: date | None,
    date_to: date | None,
) -> str:
    if date_from and date_to:
        return f"{format_report_date(date_from)} — {format_report_date(date_to)}"
    if date_from:
        return f"Dari {format_report_date(date_from)}"
    if date_to:
        return f"Sampai {format_report_date(date_to)}"
    return "Semua Periode"


def date_range_filename_suffix(
    date_from: date | None,
    date_to: date | None,
) -> str:
    if date_from and date_to:
        return f"_{date_from.isoformat()}_sd_{date_to.isoformat()}"
    if date_from:
        return f"_dari_{date_from.isoformat()}"
    if date_to:
        return f"_sampai_{date_to.isoformat()}"
    return ""


def row_height_for_text(
    text,
    *,
    chars_per_line: int = 24,
    min_height: float = 20,
    max_height: float = 120,
    line_px: float = 14,
    padding: float = 8,
) -> float:
    if text is None or text == "":
        return min_height
    lines = 0
    for part in str(text).split("\n"):
        part_len = len(part)
        lines += max(1, (part_len + chars_per_line - 1) // chars_per_line)
    return min(max_height, max(min_height, padding + lines * line_px))


def auto_column_width(
    header: str,
    values,
    *,
    min_width: float,
    max_width: float = 48,
    extra: int = 2,
) -> float:
    max_len = max(len(line) for line in header.split("\n"))
    for value in values:
        if value is None:
            continue
        text = str(value)
        for line in text.split("\n"):
            max_len = max(max_len, len(line))
    return min(max_width, max(min_width, max_len + extra))


def apply_merged_row_style(
    ws: Worksheet,
    row: int,
    start_col: int,
    end_col: int,
    *,
    fill: PatternFill | None = None,
    border: Border | None = None,
    alignment: Alignment | None = None,
    font: Font | None = None,
) -> None:
    for col in range(start_col, end_col + 1):
        cell = ws.cell(row=row, column=col)
        if fill is not None:
            cell.fill = fill
        if border is not None:
            cell.border = border
        if alignment is not None and col == start_col:
            cell.alignment = alignment
        if font is not None and col == start_col:
            cell.font = font


def configure_print_layout(
    ws: Worksheet,
    *,
    header_row: int,
    landscape: bool = True,
    paper_size: int | None = None,
) -> None:
    ws.print_title_rows = f"{header_row}:{header_row}"
    ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE if landscape else ws.ORIENTATION_PORTRAIT
    ws.page_setup.fitToPage = False
    ws.page_setup.fitToWidth = 0
    ws.page_setup.fitToHeight = 0
    if paper_size is not None:
        ws.page_setup.paperSize = paper_size
    ws.page_margins.left = 0.4
    ws.page_margins.right = 0.4
    ws.page_margins.top = 0.5
    ws.page_margins.bottom = 0.5


def finalize_column_widths(
    ws: Worksheet,
    columns: list[tuple[str, float]],
    *,
    header_row: int,
    data_start: int,
    data_end: int,
    overrides: dict[int, float] | None = None,
) -> None:
    for col_idx, (header, min_width) in enumerate(columns, start=1):
        values = [
            ws.cell(row=r, column=col_idx).value
            for r in range(data_start, data_end + 1)
        ]
        width = auto_column_width(header, values, min_width=min_width)
        if overrides and col_idx in overrides:
            width = max(width, overrides[col_idx])
        ws.column_dimensions[get_column_letter(col_idx)].width = width


def set_row_heights_for_wrapped_rows(
    ws: Worksheet,
    row: int,
    wrapped_columns: list[tuple[int, float]],
    *,
    min_height: float = 20,
) -> None:
    height = min_height
    for col_idx, col_width in wrapped_columns:
        value = ws.cell(row=row, column=col_idx).value
        if value is None:
            continue
        height = max(
            height,
            row_height_for_text(value, chars_per_line=max(int(col_width * 1.05), 10)),
        )
    ws.row_dimensions[row].height = height
