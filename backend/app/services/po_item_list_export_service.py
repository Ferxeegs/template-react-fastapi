"""
Excel export for Purchase Order items list.
"""
from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from typing import Optional
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.schemas.procurement import PurchaseOrderItemListRead
from app.services.excel_export_utils import (
    apply_merged_row_style,
    configure_print_layout,
    date_range_filename_suffix,
    finalize_column_widths,
    format_date_range_label,
    row_height_for_text,
    set_row_heights_for_wrapped_rows,
)

WIB = ZoneInfo("Asia/Jakarta")

QTY_FORMAT_INT = "#,##0"
QTY_FORMAT_DEC = "#,##0.####"
IDR_FORMAT = '"Rp" #,##0;"Rp" -#,##0'

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
TITLE_FONT = Font(name="Calibri", size=18, bold=True, color="1F4E79")
SUBTITLE_FONT = Font(name="Calibri", size=11, color="6B7280")
META_LABEL_FONT = Font(name="Calibri", size=10, bold=True, color="374151")
META_VALUE_FONT = Font(name="Calibri", size=10, color="111827")
DATA_FONT = Font(name="Calibri", size=10, color="111827")
TOTAL_FONT = Font(name="Calibri", size=10, bold=True, color="111827")
TOTAL_FILL = PatternFill("solid", fgColor="E5E7EB")
META_FILL = PatternFill("solid", fgColor="F8FAFC")
ALT_ROW_FILL = PatternFill("solid", fgColor="F9FAFB")

THIN = Side(style="thin", color="D1D5DB")
MEDIUM = Side(style="medium", color="9CA3AF")

THIN_BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
TOTAL_TOP_BORDER = Border(left=THIN, right=THIN, top=MEDIUM, bottom=THIN)

COLUMNS = [
    ("No.", 6),
    ("No. PO", 18),
    ("No. PR", 18),
    ("Vendor", 26),
    ("Nama Item", 34),
    ("Kategori", 18),
    ("Qty", 10),
    ("Satuan", 10),
    ("Harga\nSatuan (Rp)", 18),
    ("Total\nItem (Rp)", 18),
    ("Status\nPengiriman", 16),
    ("Status\nPembayaran", 16),
    ("Tanggal\nPO", 18),
]

LABEL_COL_WIDTH = 22
QTY_COL = 7
PRICE_COL = 9
TOTAL_COL = 10
WRAP_COLS = (4, 5, 6)


def po_item_list_filename(
    date_from: date | None = None,
    date_to: date | None = None,
) -> str:
    stamp = datetime.now(WIB).strftime("%Y%m%d")
    return f"laporan-item-po{date_range_filename_suffix(date_from, date_to)}_{stamp}.xlsx"


def _decimal(value) -> Decimal:
    if value is None:
        return Decimal("0")
    return Decimal(str(value))


def _is_whole_qty(value: Decimal) -> bool:
    normalized = value.quantize(Decimal("0.0001"))
    return normalized == normalized.to_integral_value()


def _qty_number_format(value: Decimal) -> str:
    return QTY_FORMAT_INT if _is_whole_qty(value) else QTY_FORMAT_DEC


def _qty_value(value: Decimal) -> float | int:
    normalized = value.quantize(Decimal("0.0001"))
    if normalized == normalized.to_integral_value():
        return int(normalized)
    return float(normalized)


def _idr_value(value: Decimal) -> int:
    return int(value.quantize(Decimal("1")))


def generate_po_item_list_excel(
    items: list[PurchaseOrderItemListRead],
    *,
    generated_by: str,
    search: Optional[str] = None,
    delivery_status: Optional[str] = None,
    payment_status: Optional[str] = None,
    category_name: Optional[str] = None,
    date_from: date | None = None,
    date_to: date | None = None,
) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Item PO"

    last_col = len(COLUMNS)
    last_col_letter = get_column_letter(last_col)

    ws.merge_cells(f"A1:{last_col_letter}1")
    title = ws["A1"]
    title.value = "LAPORAN ITEM PURCHASE ORDER"
    title.font = TITLE_FONT
    title.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 34

    ws.merge_cells(f"A2:{last_col_letter}2")
    subtitle = ws["A2"]
    subtitle.value = "Purchasing Go — Daftar Item Purchase Order"
    subtitle.font = SUBTITLE_FONT
    subtitle.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[2].height = 24

    meta_rows: list[tuple[str, str]] = [
        ("Periode Laporan", format_date_range_label(date_from, date_to)),
        ("Pencarian", search or "—"),
    ]
    if delivery_status:
        meta_rows.append(("Status Pengiriman", delivery_status))
    if payment_status:
        meta_rows.append(("Status Pembayaran", payment_status))
    if category_name:
        meta_rows.append(("Kategori", category_name))
    meta_rows.extend([
        ("Waktu Cetak", datetime.now(WIB).strftime("%d/%m/%Y %H:%M WIB")),
        ("Petugas / User", generated_by),
        ("Jumlah Item", str(len(items))),
    ])

    meta_start = 4
    meta_end = meta_start + len(meta_rows) - 1
    ws.column_dimensions["A"].width = LABEL_COL_WIDTH

    for offset, (label, value) in enumerate(meta_rows):
        row_num = meta_start + offset
        label_cell = ws.cell(row=row_num, column=1, value=label)
        label_cell.font = META_LABEL_FONT
        label_cell.fill = META_FILL
        label_cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        label_cell.border = THIN_BORDER

        ws.merge_cells(start_row=row_num, start_column=2, end_row=row_num, end_column=last_col)
        value_cell = ws.cell(row=row_num, column=2, value=value)
        value_cell.font = META_VALUE_FONT
        value_cell.alignment = Alignment(horizontal="left", vertical="center", indent=1, wrap_text=True)
        apply_merged_row_style(
            ws,
            row_num,
            2,
            last_col,
            fill=META_FILL,
            border=THIN_BORDER,
        )
        ws.row_dimensions[row_num].height = row_height_for_text(value, chars_per_line=90, min_height=22)

    header_row = meta_end + 2
    for col_idx, (header, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=header_row, column=col_idx, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.border = THIN_BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    ws.row_dimensions[header_row].height = 40

    data_start = header_row + 1
    total_amount = Decimal("0")

    if not items:
        ws.merge_cells(start_row=data_start, start_column=1, end_row=data_start, end_column=last_col)
        note = ws.cell(row=data_start, column=1, value="Tidak ada data item purchase order sesuai filter.")
        note.font = Font(name="Calibri", size=10, italic=True, color="6B7280")
        note.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        apply_merged_row_style(ws, data_start, 1, last_col, border=THIN_BORDER)
        ws.row_dimensions[data_start].height = 28
        last_data_row = data_start
    else:
        for idx, item in enumerate(items):
            row_num = data_start + idx
            row_fill = ALT_ROW_FILL if idx % 2 == 1 else None
            qty = _decimal(item.real_qty)
            unit_price = _decimal(item.real_price)
            line_total = _decimal(item.real_total)
            total_amount += line_total

            values = [
                idx + 1,
                item.po_number,
                item.pr_number or "-",
                item.vendor_name or "-",
                item.product_name,
                item.category_name or "-",
                _qty_value(qty),
                item.uom or "-",
                _idr_value(unit_price),
                _idr_value(line_total),
                item.delivery_status.name if item.delivery_status else "-",
                item.payment_status.name if item.payment_status else "-",
                item.po_created_at.astimezone(WIB).strftime("%d/%m/%Y %H:%M") if item.po_created_at else "-",
            ]
            alignments = [
                "center", "left", "left", "left", "left", "left",
                "right", "center", "right", "right", "center", "center", "center",
            ]

            for col_idx, (value, align) in enumerate(zip(values, alignments), start=1):
                cell = ws.cell(row=row_num, column=col_idx, value=value)
                cell.font = DATA_FONT
                cell.alignment = Alignment(
                    horizontal=align,
                    vertical="center",
                    wrap_text=col_idx in WRAP_COLS or col_idx in (2, 3),
                )
                cell.border = THIN_BORDER
                if row_fill:
                    cell.fill = row_fill
                if col_idx == QTY_COL:
                    cell.number_format = _qty_number_format(qty)
                elif col_idx in (PRICE_COL, TOTAL_COL):
                    cell.number_format = IDR_FORMAT

            set_row_heights_for_wrapped_rows(
                ws,
                row_num,
                [(5, COLUMNS[4][1]), (4, COLUMNS[3][1]), (6, COLUMNS[5][1])],
                min_height=22,
            )

        last_data_row = data_start + len(items) - 1

    total_row = last_data_row + 1
    ws.merge_cells(start_row=total_row, start_column=1, end_row=total_row, end_column=TOTAL_COL - 1)
    total_label = ws.cell(row=total_row, column=1, value="JUMLAH")
    total_label.font = TOTAL_FONT
    total_label.fill = TOTAL_FILL
    total_label.alignment = Alignment(horizontal="right", vertical="center", indent=1)
    total_label.border = TOTAL_TOP_BORDER
    apply_merged_row_style(ws, total_row, 1, TOTAL_COL - 1, fill=TOTAL_FILL, border=TOTAL_TOP_BORDER)

    total_cell = ws.cell(row=total_row, column=TOTAL_COL, value=_idr_value(total_amount))
    total_cell.font = TOTAL_FONT
    total_cell.fill = TOTAL_FILL
    total_cell.border = TOTAL_TOP_BORDER
    total_cell.alignment = Alignment(horizontal="right", vertical="center")
    total_cell.number_format = IDR_FORMAT

    for col_idx in range(TOTAL_COL + 1, last_col + 1):
        empty = ws.cell(row=total_row, column=col_idx, value="")
        empty.fill = TOTAL_FILL
        empty.border = TOTAL_TOP_BORDER
    ws.row_dimensions[total_row].height = 24

    if items:
        finalize_column_widths(
            ws,
            COLUMNS,
            header_row=header_row,
            data_start=data_start,
            data_end=last_data_row,
            overrides={PRICE_COL: 18, TOTAL_COL: 18, QTY_COL: 10},
        )

    footer_row = total_row + 2
    ws.merge_cells(start_row=footer_row, start_column=1, end_row=footer_row + 1, end_column=last_col)
    footer = ws.cell(
        row=footer_row,
        column=1,
        value=(
            "Catatan: Data diekspor sesuai filter yang dipilih. "
            "Total dihitung dari jumlah nilai item yang ditampilkan. "
            "Dokumen ini dihasilkan otomatis oleh sistem Purchasing Go."
        ),
    )
    footer.font = Font(name="Calibri", size=9, italic=True, color="9CA3AF")
    footer.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
    ws.row_dimensions[footer_row].height = row_height_for_text(
        footer.value,
        chars_per_line=100,
        min_height=40,
        max_height=80,
    )
    ws.row_dimensions[footer_row + 1].height = 8

    ws.freeze_panes = ws.cell(row=data_start, column=1).coordinate
    ws.sheet_view.showGridLines = False
    configure_print_layout(ws, header_row=header_row, landscape=True)

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
