"""
Excel export for daily stock summary report.
"""
from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from typing import Optional
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from app.models.purchasing import Product
from app.schemas.inventory import StockDailySummaryRead
from app.services.excel_export_utils import (
    apply_merged_row_style,
    configure_print_layout,
    date_range_filename_suffix,
    finalize_column_widths,
    format_date_range_label,
    row_height_for_text,
    set_row_heights_for_wrapped_rows,
)

WIB = ZoneInfo("Asia/Jakarta")

# Format US disimpan di xlsx; Excel Indonesia menampilkan ribuan "." dan desimal ",".
# Bilangan bulat: tanpa bagian desimal sama sekali agar tidak muncul ",0".
QTY_FORMAT_INT = "#,##0"
QTY_FORMAT_DEC = "#,##0.####"

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
TITLE_FONT = Font(name="Calibri", size=18, bold=True, color="1F4E79")
SUBTITLE_FONT = Font(name="Calibri", size=11, color="6B7280")
META_LABEL_FONT = Font(name="Calibri", size=10, bold=True, color="374151")
META_VALUE_FONT = Font(name="Calibri", size=10, color="111827")
DATA_FONT = Font(name="Calibri", size=10, color="111827")
TOTAL_FONT = Font(name="Calibri", size=10, bold=True, color="111827")
TOTAL_FILL = PatternFill("solid", fgColor="E5E7EB")
META_FILL = PatternFill("solid", fgColor="F8FAFC")
ALT_ROW_FILL = PatternFill("solid", fgColor="F9FAFB")

THIN = Side(style="thin", color="D1D5DB")
MEDIUM = Side(style="medium", color="9CA3AF")

THIN_BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
TOTAL_TOP_BORDER = Border(
    left=THIN, right=THIN, top=MEDIUM, bottom=THIN,
)

COLUMNS = [
    ("No.", 6),
    ("Nama Produk", 38),
    ("Kategori", 26),
    ("Unit Kerja", 24),
    ("Satuan", 14),
    ("Stok\nAwal", 14),
    ("Masuk\n(+)", 14),
    ("Keluar\n(-)", 14),
    ("Stok\nAkhir", 14),
]

LABEL_COL_WIDTH = 22
QTY_COL_START = 6
QTY_COL_END = 9
QTY_COL_ATTRS = ("opening_qty", "total_in", "total_out", "closing_qty")


def _format_report_date(value: date) -> str:
    months = [
        "Januari", "Februari", "Maret", "April", "Mei", "Juni",
        "Juli", "Agustus", "September", "Oktober", "November", "Desember",
    ]
    return f"{value.day:02d} {months[value.month - 1]} {value.year}"


def _decimal(value) -> Decimal:
    if value is None:
        return Decimal("0")
    return Decimal(str(value))


def _is_whole_qty(value: Decimal) -> bool:
    normalized = value.quantize(Decimal("0.0001"))
    return normalized == normalized.to_integral_value()


def _qty_number_format(value: Decimal) -> str:
    return QTY_FORMAT_INT if _is_whole_qty(value) else QTY_FORMAT_DEC


def _column_qty_format(items, attr: str) -> str:
    for item in items:
        if not _is_whole_qty(_decimal(getattr(item, attr))):
            return QTY_FORMAT_DEC
    return QTY_FORMAT_INT


def _qty_float(value: Decimal) -> float | int:
    normalized = value.quantize(Decimal("0.0001"))
    if normalized == normalized.to_integral_value():
        return int(normalized)
    return float(normalized)


def _uom_for_product(db: Session, product_id: str) -> str:
    product = db.query(Product).filter(Product.id == product_id).first()
    if not product:
        return "-"
    db.refresh(product, ["uom"])
    if product.uom:
        return product.uom.shortname or product.uom.name or "-"
    return "-"


def daily_summary_filename(summary: StockDailySummaryRead) -> str:
    date_from = summary.date_from or summary.date
    date_to = summary.date_to or summary.date
    stamp = datetime.now(WIB).strftime("%Y%m%d")
    return f"laporan-stok{date_range_filename_suffix(date_from, date_to)}_{stamp}.xlsx"


def generate_daily_stock_excel(
    db: Session,
    summary: StockDailySummaryRead,
    *,
    generated_by: Optional[str] = None,
) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Laporan Harian"

    last_col = len(COLUMNS)
    last_col_letter = get_column_letter(last_col)

    # --- Title block ---
    ws.merge_cells(f"A1:{last_col_letter}1")
    title = ws["A1"]
    title.value = "LAPORAN PERGERAKAN STOK"
    title.font = TITLE_FONT
    title.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 34

    ws.merge_cells(f"A2:{last_col_letter}2")
    subtitle = ws["A2"]
    subtitle.value = "Purchasing Go — Ringkasan Pergerakan Stok"
    subtitle.font = SUBTITLE_FONT
    subtitle.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[2].height = 24

    # --- Metadata panel ---
    meta_rows = [
        ("Periode Laporan", format_date_range_label(summary.date_from or summary.date, summary.date_to or summary.date)),
        ("Unit Kerja", summary.unit_name or "Semua Unit"),
        ("Kategori Produk", summary.category_name or "Semua Kategori"),
        ("Waktu Cetak", datetime.now(WIB).strftime("%d/%m/%Y %H:%M WIB")),
    ]
    if generated_by:
        meta_rows.append(("Petugas / User", generated_by))

    meta_start = 4
    meta_end = meta_start + len(meta_rows) - 1

    for offset, (label, value) in enumerate(meta_rows):
        row = meta_start + offset
        
        # Merge A-B for the label
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
        label_text = f"{label}:" if not label.endswith(":") else label
        label_cell = ws.cell(row=row, column=1, value=label_text)
        label_cell.font = META_LABEL_FONT
        label_cell.alignment = Alignment(horizontal="left", vertical="center")

        # Merge C-last_col for the value
        ws.merge_cells(start_row=row, start_column=3, end_row=row, end_column=last_col)
        value_cell = ws.cell(row=row, column=3, value=value)
        value_cell.font = META_VALUE_FONT
        value_cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        
        apply_merged_row_style(
            ws,
            row,
            1,
            last_col,
            fill=None,
            border=None,
        )
        ws.row_dimensions[row].height = row_height_for_text(value, chars_per_line=70, min_height=20)

    # --- Table header ---
    header_row = meta_end + 2
    for col_idx, (header, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=header_row, column=col_idx, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.border = THIN_BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    ws.row_dimensions[header_row].height = 40

    data_start = header_row + 1
    items = summary.items or []

    if not items:
        ws.merge_cells(start_row=data_start, start_column=1, end_row=data_start, end_column=last_col)
        note = ws.cell(row=data_start, column=1, value="Tidak ada pergerakan stok pada periode ini.")
        note.font = Font(name="Calibri", size=10, italic=True, color="6B7280")
        note.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        apply_merged_row_style(ws, data_start, 1, last_col, border=THIN_BORDER)
        ws.row_dimensions[data_start].height = 28
        total_row = data_start + 1
        last_data_row = data_start
    else:
        for index, item in enumerate(items, start=1):
            row = data_start + index - 1
            uom = _uom_for_product(db, item.product_id)
            opening = _decimal(item.opening_qty)
            total_in = _decimal(item.total_in)
            total_out = _decimal(item.total_out)
            closing = _decimal(item.closing_qty)

            row_fill = ALT_ROW_FILL if index % 2 == 0 else None

            no_cell = ws.cell(row=row, column=1, value=index)
            no_cell.font = DATA_FONT
            no_cell.alignment = Alignment(horizontal="center", vertical="center")
            no_cell.border = THIN_BORDER
            if row_fill:
                no_cell.fill = row_fill

            text_values = [item.product_name, item.category_name, item.unit_name, uom]
            for col_offset, text in enumerate(text_values, start=2):
                cell = ws.cell(row=row, column=col_offset, value=text)
                cell.font = DATA_FONT
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
                cell.border = THIN_BORDER
                if row_fill:
                    cell.fill = row_fill

            qty_values = [opening, total_in, total_out, closing]
            for col_offset, qty in enumerate(qty_values, start=QTY_COL_START):
                cell = ws.cell(row=row, column=col_offset, value=_qty_float(qty))
                cell.font = DATA_FONT
                cell.alignment = Alignment(horizontal="right", vertical="center")
                cell.number_format = _qty_number_format(qty)
                cell.border = THIN_BORDER
                if row_fill:
                    cell.fill = row_fill

            set_row_heights_for_wrapped_rows(
                ws,
                row,
                [(2, COLUMNS[1][1]), (3, COLUMNS[2][1]), (4, COLUMNS[3][1])],
                min_height=22,
            )

        total_row = data_start + len(items)
        last_data_row = total_row - 1

    # --- Total row ---
    ws.merge_cells(start_row=total_row, start_column=1, end_row=total_row, end_column=5)
    total_label = ws.cell(row=total_row, column=1, value="JUMLAH")
    total_label.font = TOTAL_FONT
    total_label.fill = TOTAL_FILL
    total_label.alignment = Alignment(horizontal="right", vertical="center", indent=1)
    total_label.border = TOTAL_TOP_BORDER
    apply_merged_row_style(ws, total_row, 1, 5, fill=TOTAL_FILL, border=TOTAL_TOP_BORDER)

    if items:
        first_data_row = data_start
        last_data_row = data_start + len(items) - 1
        for offset, col in enumerate(range(QTY_COL_START, QTY_COL_END + 1)):
            col_letter = get_column_letter(col)
            attr = QTY_COL_ATTRS[offset]
            cell = ws.cell(
                row=total_row,
                column=col,
                value=f"=SUM({col_letter}{first_data_row}:{col_letter}{last_data_row})",
            )
            cell.font = TOTAL_FONT
            cell.fill = TOTAL_FILL
            cell.border = TOTAL_TOP_BORDER
            cell.alignment = Alignment(horizontal="right", vertical="center")
            cell.number_format = _column_qty_format(items, attr)
    else:
        for col in range(QTY_COL_START, QTY_COL_END + 1):
            cell = ws.cell(row=total_row, column=col, value=0)
            cell.font = TOTAL_FONT
            cell.fill = TOTAL_FILL
            cell.border = TOTAL_TOP_BORDER
            cell.alignment = Alignment(horizontal="right", vertical="center")
            cell.number_format = QTY_FORMAT_INT

    ws.row_dimensions[total_row].height = 24

    if items:
        finalize_column_widths(
            ws,
            COLUMNS,
            header_row=header_row,
            data_start=data_start,
            data_end=last_data_row,
            overrides={col: 14 for col in range(QTY_COL_START, QTY_COL_END + 1)},
        )

    # --- Footer ---
    footer_row = total_row + 2
    ws.merge_cells(start_row=footer_row, start_column=1, end_row=footer_row + 1, end_column=last_col)
    footer = ws.cell(
        row=footer_row,
        column=1,
        value=(
            "Catatan: Stok Akhir = Stok Awal + Masuk − Keluar. "
            "Dokumen ini dihasilkan otomatis oleh sistem Purchasing Go dan sah tanpa tanda tangan basah."
        ),
    )
    footer.font = Font(name="Calibri", size=9, italic=True, color="9CA3AF")
    footer.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
    ws.row_dimensions[footer_row].height = row_height_for_text(
        footer.value,
        chars_per_line=90,
        min_height=40,
        max_height=80,
    )
    ws.row_dimensions[footer_row + 1].height = 8

    ws.freeze_panes = ws.cell(row=data_start, column=1).coordinate
    ws.sheet_view.showGridLines = True
    configure_print_layout(ws, header_row=header_row, landscape=True)

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
